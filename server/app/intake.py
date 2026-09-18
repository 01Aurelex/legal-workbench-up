# -*- coding: utf-8 -*-
"""接案笔录（v0.9.8）

全本地流程：浏览器录音/上传音频 → faster-whisper 分段转写 →
本地大模型（或启发式）区分律师/当事人角色、提取接待要素 →
按律协接待谈话笔录规范生成可编辑 Word，并把录音、转写稿、笔录统一归档。

规范依据（公开行业指引）：
- 中华全国律师协会《律师执业行为规范》《律师办理民事诉讼案件规范》关于接待登记、
  谈话笔录、风险告知的要求；
- 地方律协业务操作指引中《来访登记表》《接待谈话笔录》通行要素：
  时间地点人物、当事人身份信息、案情事实、现有证据、诉求、
  律师分析与风险告知、需补充材料、当事人阅后签名。
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path

from .config import VAULT_DIR
from . import db, media, llm, archive, vault

DOC_DIR = VAULT_DIR / "文书" / "接案笔录"
AUDIO_DIR = VAULT_DIR / "材料" / "接案笔录录音"
TRANSCRIPT_DIR = DOC_DIR / "转写稿"

ROLE_LAWYER = "lawyer"
ROLE_CLIENT = "client"
ROLE_OTHER = "other"
ROLE_NAME = {ROLE_LAWYER: "律师", ROLE_CLIENT: "当事人", ROLE_OTHER: "其他"}

# 笔录中固定的律师告知事项（可在界面对应生成结果里于 Word 中继续编辑）
STANDARD_NOTICES = [
    "1. 诉讼/仲裁风险告知：案件结果受事实证据、法律适用及裁判机关自由裁量等因素影响，"
    "律师依法提供法律服务，不对案件结果作胜诉承诺。",
    "2. 举证责任告知：当事人对自己提出的主张有责任提供证据；应如实陈述并及时提交证据，"
    "逾期举证或证据不足可能承担不利后果。",
    "3. 程序与时效告知：诉讼时效、上诉期、举证期限、申请保全及申请执行均有法定期间，"
    "因当事人延误造成的后果由当事人自行承担。",
    "4. 保全与执行风险告知：财产保全需提供担保并缴纳申请费；胜诉后债权实现取决于"
    "被执行人履行能力与财产状况，存在执行不能的风险。",
    "5. 收费告知：律师服务费依据《律师服务收费管理办法》及本所收费标准，"
    "在签订《委托代理合同》时另行明确；办案中发生的诉讼费、保全费、鉴定费等由当事人承担。",
    "6. 利益冲突审查：本所将依法进行利益冲突检索，如发现依法不得代理的情形将及时告知。",
    "7. 保密义务：律师对执业活动中知悉的当事人信息和案件情况依法承担保密义务。",
]

META_FIELDS = ["client", "gender", "id_no", "phone", "address", "work_unit",
               "cause", "claims", "facts", "evidence", "materials", "risk_notes",
               "lawyer", "recorder", "location"]


# ----------------------------- 基础工具 -----------------------------
def _safe_name(name: str, fallback: str = "未命名") -> str:
    name = (name or "").strip() or fallback
    name = re.sub(r'[\\/:*?"<>|\r\n\t]', "_", name)
    return name[:80]


def _parse_dt(s: str) -> datetime | None:
    if not s:
        return None
    for fmt in ("%Y-%m-%dT%H:%M:%S.%f", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S",
                "%Y-%m-%d %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt)
        except ValueError:
            continue
    try:
        return datetime.fromisoformat(s)
    except Exception:
        return None


def _extract_json(text: str):
    """从模型输出中稳健提取 JSON（对象或数组）。"""
    if not text:
        return None
    fence = re.search(r"```(?:json)?\s*(.*?)```", text, re.S)
    if fence:
        text = fence.group(1)
    for op, cl in (("[", "]"), ("{", "}")):
        i, j = text.find(op), text.rfind(cl)
        if i >= 0 and j > i:
            frag = text[i:j + 1]
            try:
                return json.loads(frag)
            except Exception:
                pass
    return None


# ----------------------------- 角色区分 -----------------------------
# 强疑问特征：问候语（您好）不算提问，避免把当事人自我介绍误判为律师
_QUESTION_WORDS = ("请问", "是否", "有没有", "是不是", "能否", "哪些", "什么",
                   "怎么", "为什么", "何时", "哪里", "多少", "几年", "几次",
                   "明白吗", "清楚吗", "好不好", "对吗", "是吗")


def _is_question(t: str) -> bool:
    t = (t or "").strip()
    if not t:
        return False
    if t.endswith("？") or t.endswith("?") or "吗？" in t or "吗?" in t:
        return True
    return any(w in t for w in _QUESTION_WORDS) and len(t) <= 120


def heuristic_roles(segments: list[dict]) -> list[str]:
    """无本地模型时的兜底角色判断：提问归律师、陈述归当事人，并按问答交替修正。"""
    roles = []
    prev = ROLE_LAWYER
    for idx, seg in enumerate(segments):
        t = seg.get("text", "")
        is_q = _is_question(t)
        if idx == 0:
            # 首段：提问或简短开场归律师，否则归当事人
            role = ROLE_LAWYER if (is_q or len(t) <= 40) else ROLE_CLIENT
        elif is_q:
            role = ROLE_LAWYER
        else:
            # 陈述段：与上一说话人相反（问答交替），连续陈述保持当事人
            role = ROLE_CLIENT if prev in (ROLE_LAWYER, ROLE_OTHER) else ROLE_CLIENT
        roles.append(role)
        prev = role
    return roles


def attribute_roles(segments: list[dict]) -> tuple[list[str], str]:
    """返回 (角色列表, 引擎来源 llm/heuristic)。"""
    if not segments:
        return [], "heuristic"
    numbered = "\n".join(f"{i + 1}. {s['text']}" for i, s in enumerate(segments))
    prompt = (
        "你是中国律师事务所的接待助理。下面是律师初次接待当事人的录音转写片段，"
        "已按时间顺序编号。请判断每一段的说话人角色：\n"
        "- lawyer：接待律师（提问、引导陈述、释法、告知风险与程序）\n"
        "- client：当事人/委托人（陈述事实、回答问题、表达诉求）\n"
        "- other：其他在场人员\n"
        "只输出一个 JSON 字符串数组，长度与片段数完全一致，元素只能是 "
        "\"lawyer\"/\"client\"/\"other\"，不要输出任何其他文字。\n\n"
        f"片段：\n{numbered}")
    try:
        out = llm.chat([{"role": "user", "content": prompt}], timeout=90)
        arr = _extract_json(out)
        if isinstance(arr, list) and len(arr) == len(segments):
            roles = []
            for x in arr:
                x = str(x).strip().lower()
                roles.append(x if x in ROLE_NAME else ROLE_OTHER)
            return roles, "llm"
    except Exception:
        pass
    return heuristic_roles(segments), "heuristic"


# ----------------------------- 要素提取 -----------------------------
def extract_facts(segments: list[dict], roles: list[str]) -> tuple[dict, str]:
    dialogue = "\n".join(
        f"{ROLE_NAME.get(roles[i] if i < len(roles) else ROLE_OTHER)}：{s['text']}"
        for i, s in enumerate(segments))
    blank = {k: "" for k in META_FIELDS}
    prompt = (
        "下面是一份律师接待当事人的谈话记录。请提取并归纳接待笔录要素，"
        "只输出 JSON 对象，键固定为：client(当事人姓名或单位名称), gender(性别), "
        "id_no(身份证号/统一社会信用代码), phone(联系电话), address(住址/住所地), "
        "work_unit(工作单位), cause(案由或咨询事项), claims(当事人诉求，分号分隔), "
        "facts(主要事实经过，200字以内), evidence(现有证据清单，分号分隔), "
        "materials(尚需补充的材料，分号分隔), risk_notes(律师提示的主要风险)。"
        "无法确定的键值留空字符串，不要编造。\n\n谈话记录：\n" + dialogue)
    try:
        out = llm.chat([{"role": "user", "content": prompt}], timeout=90)
        obj = _extract_json(out)
        if isinstance(obj, dict):
            return {**blank, **{k: str(obj.get(k, "") or "") for k in META_FIELDS}}, "llm"
    except Exception:
        pass
    return blank, "heuristic"


# ----------------------------- 转写建档 -----------------------------
def create_from_audio(audio_bytes: bytes, filename: str, started_at: str = "",
                      client: str = "", lawyer: str = "", recorder: str = "",
                      location: str = "") -> dict:
    """保存录音 → 分段转写 → 角色区分 → 要素提取 → 写 intake_records。"""
    suffix = Path(filename or "audio.webm").suffix.lower() or ".webm"
    if suffix not in (".webm", ".wav", ".mp3", ".m4a", ".ogg", ".opus", ".mp4", ".aac", ".flac"):
        suffix = ".webm"

    started = _parse_dt(started_at) or datetime.now()
    stamp = started.strftime("%Y%m%d_%H%M%S")

    # 1) 录音原件归档（保留在本地 vault）
    AUDIO_DIR.mkdir(parents=True, exist_ok=True)
    audio_name = f"{stamp}_{_safe_name(client or '未命名当事人')}{suffix}"
    audio_path = AUDIO_DIR / audio_name
    n = 1
    while audio_path.exists():
        audio_path = AUDIO_DIR / f"{stamp}_{_safe_name(client or '未命名当事人')}_{n}{suffix}"
        n += 1
    audio_path.write_bytes(audio_bytes)

    # 2) 分段转写
    tr = media.transcribe_segments(audio_bytes, suffix=suffix)
    if not tr.get("ok"):
        # 转写失败仍保留录音与记录，允许稍后重试
        segments = []
        duration = 0.0
        engine = ""
        err = tr.get("error", "转写失败")
    else:
        segments = tr["segments"]
        duration = float(tr.get("duration") or 0)
        engine = tr.get("engine", "")
        err = ""

    # 3) 角色区分 + 要素提取
    roles, role_engine = attribute_roles(segments)
    meta, fact_engine = extract_facts(segments, roles)
    meta.update({"client": client or meta.get("client", ""),
                 "lawyer": lawyer, "recorder": recorder, "location": location})
    turns = [{"role": roles[i] if i < len(roles) else ROLE_OTHER,
              "start": s["start"], "end": s["end"], "text": s["text"]}
             for i, s in enumerate(segments)]

    rec_id = db.insert(
        "intake_records",
        {"created": db.now(), "started_at": started.strftime("%Y-%m-%d %H:%M"),
         "ended_at": datetime.now().strftime("%Y-%m-%d %H:%M"),
         "client": meta.get("client", ""), "lawyer": lawyer, "recorder": recorder,
         "location": location, "duration": duration,
         "audio_path": str(audio_path.relative_to(VAULT_DIR)).replace("\\", "/"),
         "status": "transcribed", "summary": meta.get("cause", ""),
         "turns_json": json.dumps(turns, ensure_ascii=False),
         "meta_json": json.dumps(meta, ensure_ascii=False)})
    db.audit("接案笔录转写", f"{meta.get('client') or client or '未命名'} · {len(turns)} 段")
    return {"ok": not bool(err), "id": rec_id, "engine": engine,
            "role_engine": role_engine, "fact_engine": fact_engine,
            "duration": duration, "turns": turns, "meta": meta,
            "audio_path": str(audio_path.relative_to(VAULT_DIR)).replace("\\", "/"),
            "error": err}


# ----------------------------- Word 生成 -----------------------------
def _set_doc_default_font(doc):
    from docx.oxml.ns import qn
    style = doc.styles["Normal"]
    style.font.name = "FangSong"
    style.font.size = None
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "仿宋")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def _para(doc, text="", *, size=12, bold=False, align=None, eastasia="仿宋",
          space_after=6, indent=False):
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if indent:
        pf.first_line_indent = Pt(size * 2)
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.name = "Times New Roman"
    from docx.oxml.ns import qn
    run._element.rPr.rFonts.set(qn("w:eastAsia"), eastasia)
    return p


def _merge_turns(turns: list[dict]) -> list[dict]:
    """合并相邻同角色片段。"""
    merged = []
    for t in turns:
        text = (t.get("text") or "").strip()
        if not text:
            continue
        role = t.get("role", ROLE_OTHER)
        if merged and merged[-1]["role"] == role:
            merged[-1]["text"] += text
        else:
            merged.append({"role": role, "text": text})
    return merged


def build_docx(record: dict, meta: dict, turns: list[dict]) -> Path:
    from docx import Document
    from docx.shared import Pt, Cm
    doc = Document()
    _set_doc_default_font(doc)
    for sec in doc.sections:
        sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
        sec.left_margin = Cm(2.8); sec.right_margin = Cm(2.6)

    _para(doc, "接待当事人谈话笔录", size=22, bold=True, align="center",
          eastasia="黑体", space_after=4)
    _para(doc, "（接案笔录）", size=12, align="center", eastasia="仿宋", space_after=12)

    # 表头信息
    started = record.get("started_at") or ""
    ended = record.get("ended_at") or ""
    _para(doc, f"时　间：{started} 至 {ended}", size=12)
    _para(doc, f"地　点：{meta.get('location','')}", size=12)
    _para(doc, f"接待律师：{meta.get('lawyer','')}　　记录人：{meta.get('recorder','')}", size=12)

    # 当事人信息
    client_line = (f"当事人：{meta.get('client','')}"
                   + (f"　性别：{meta.get('gender')}" if meta.get("gender") else "")
                   + (f"　电话：{meta.get('phone')}" if meta.get("phone") else ""))
    _para(doc, client_line, size=12)
    if meta.get("id_no"):
        _para(doc, f"证件号码：{meta.get('id_no')}", size=12)
    if meta.get("work_unit"):
        _para(doc, f"工作单位：{meta.get('work_unit')}", size=12)
    if meta.get("address"):
        _para(doc, f"住　址：{meta.get('address')}", size=12)
    if meta.get("cause"):
        _para(doc, f"咨询/委托事项：{meta.get('cause')}", size=12)

    _para(doc, "谈话内容：", size=12, bold=True, space_after=4)
    merged = _merge_turns(turns)
    if not merged:
        _para(doc, "（录音未识别出文字，可在此手工补记。）", size=12)
    for m in merged:
        prefix = {"lawyer": "问：", "client": "答："}.get(m["role"], "？：")
        _para(doc, prefix + m["text"], size=12, indent=False)

    def sec_block(title, value):
        _para(doc, title, size=12, bold=True, space_after=2)
        if value:
            for line in str(value).split(";") if False else [str(value)]:
                _para(doc, line.strip(), size=12, indent=True)
        else:
            _para(doc, "（待补充）", size=12, indent=True)

    sec_block("一、主要事实经过", meta.get("facts", ""))
    sec_block("二、现有证据", meta.get("evidence", ""))
    sec_block("三、当事人诉求", meta.get("claims", ""))

    _para(doc, "四、律师告知事项", size=12, bold=True, space_after=2)
    notices = STANDARD_NOTICES[:]
    if meta.get("risk_notes"):
        notices.append(f"8. 本案特别风险提示：{meta.get('risk_notes')}")
    for nt in notices:
        _para(doc, nt, size=12, indent=False)
    sec_block("五、需当事人补充的材料", meta.get("materials", ""))
    _para(doc, "六、委托意向：□ 当场建立委托　□ 回去商议后回复　□ 仅咨询不委托",
          size=12)

    _para(doc, "", size=12, space_after=10)
    _para(doc, "以上笔录共　　页，我已阅读（或已向我宣读），记录内容与我陈述一致。",
          size=12)
    _para(doc, "", size=12, space_after=8)
    _para(doc, "当事人（签名/捺印）：　　　　　　　　　日期：　　　年　　月　　日",
          size=12)
    _para(doc, "接待律师（签名）：　　　　　　　　　　记录人（签名）：　　　　　　　",
          size=12)

    DOC_DIR.mkdir(parents=True, exist_ok=True)
    started_dt = _parse_dt(record.get("started_at", "")) or datetime.now()
    base = f"{started_dt.strftime('%Y%m%d%H%M')}_{_safe_name(meta.get('client') or record.get('client') or '未命名当事人')}_接案笔录"
    out = DOC_DIR / f"{base}.docx"
    n = 1
    while out.exists():
        out = DOC_DIR / f"{base}_{n}.docx"; n += 1
    doc.save(out)
    return out


def _build_transcript_md(record: dict, meta: dict, turns: list[dict]) -> Path:
    TRANSCRIPT_DIR.mkdir(parents=True, exist_ok=True)
    started_dt = _parse_dt(record.get("started_at", "")) or datetime.now()
    base = f"{started_dt.strftime('%Y%m%d%H%M')}_{_safe_name(meta.get('client') or '未命名当事人')}_转写稿"
    out = TRANSCRIPT_DIR / f"{base}.md"
    n = 1
    while out.exists():
        out = TRANSCRIPT_DIR / f"{base}_{n}.md"; n += 1
    lines = ["---", "type: 接案笔录转写",
             f"client: \"{meta.get('client','')}\"",
             f"created: {db.now()}", "---", "",
             f"# 接待谈话转写稿（{record.get('started_at','')}）", "",
             f"- 接待律师：{meta.get('lawyer','')}",
             f"- 地点：{meta.get('location','')}", ""]
    for t in _merge_turns(turns):
        lines.append(f"**{ROLE_NAME.get(t['role'],'其他')}**：{t['text']}")
        lines.append("")
    out.write_text("\n".join(lines), encoding="utf-8")
    return out


def generate(record_id: int, meta: dict, turns: list[dict]) -> dict:
    """根据用户校对后的要素与对话生成 Word 笔录与转写 md。"""
    rec = db.query_one("SELECT * FROM intake_records WHERE id=?", (int(record_id),))
    if not rec:
        return {"ok": False, "error": "记录不存在"}
    meta = {**{k: "" for k in META_FIELDS}, **(meta or {})}
    docx_path = build_docx(dict(rec), meta, turns or [])
    md_path = _build_transcript_md(dict(rec), meta, turns or [])

    rel_docx = str(docx_path.relative_to(VAULT_DIR)).replace("\\", "/")
    rel_md = str(md_path.relative_to(VAULT_DIR)).replace("\\", "/")
    db.execute(
        "UPDATE intake_records SET status='done', docx_path=?, transcript_path=?, "
        "client=?, lawyer=?, recorder=?, location=?, summary=?, turns_json=?, meta_json=? WHERE id=?",
        (rel_docx, rel_md, meta.get("client", ""), meta.get("lawyer", ""),
         meta.get("recorder", ""), meta.get("location", ""), meta.get("cause", ""),
         json.dumps(turns or [], ensure_ascii=False),
         json.dumps(meta, ensure_ascii=False), int(record_id)))
    try:
        vault.scan_vault()
        archive.sync_registry()
    except Exception:
        pass
    db.audit("生成接案笔录", f"{meta.get('client','')} → {docx_path.name}")
    return {"ok": True, "docx_path": rel_docx, "transcript_path": rel_md,
            "docx_name": docx_path.name}


def list_records() -> list[dict]:
    rows = db.query("SELECT * FROM intake_records ORDER BY id DESC LIMIT 200")
    out = []
    for r in rows:
        d = dict(r)
        d.pop("turns_json", None); d.pop("meta_json", None)
        out.append(d)
    return out


def get_record(record_id: int) -> dict:
    r = db.query_one("SELECT * FROM intake_records WHERE id=?", (int(record_id),))
    if not r:
        return {"ok": False, "error": "记录不存在"}
    d = dict(r)
    try:
        d["turns"] = json.loads(d.pop("turns_json") or "[]")
    except Exception:
        d["turns"] = []
    try:
        d["meta"] = json.loads(d.pop("meta_json") or "{}")
    except Exception:
        d["meta"] = {}
    return {"ok": True, "record": d}


def retranscribe(record_id: int) -> dict:
    """对已保存录音重新转写（更换模型/首次失败后重试）。"""
    r = db.query_one("SELECT * FROM intake_records WHERE id=?", (int(record_id),))
    if not r:
        return {"ok": False, "error": "记录不存在"}
    audio = VAULT_DIR / (r["audio_path"] or "")
    if not audio.exists():
        return {"ok": False, "error": "录音文件不存在"}
    tr = media.transcribe_segments(audio.read_bytes(), suffix=audio.suffix)
    if not tr.get("ok"):
        return {"ok": False, "error": tr.get("error", "转写失败")}
    segments = tr["segments"]
    roles, role_engine = attribute_roles(segments)
    try:
        meta = json.loads(r["meta_json"] or "{}")
    except Exception:
        meta = {}
    meta2, fact_engine = extract_facts(segments, roles)
    meta2.update({k: meta.get(k, "") for k in ("client", "lawyer", "recorder", "location") if meta.get(k)})
    turns = [{"role": roles[i] if i < len(roles) else ROLE_OTHER,
              "start": s["start"], "end": s["end"], "text": s["text"]}
             for i, s in enumerate(segments)]
    db.execute("UPDATE intake_records SET turns_json=?, meta_json=?, duration=? WHERE id=?",
               (json.dumps(turns, ensure_ascii=False),
                json.dumps(meta2, ensure_ascii=False), float(tr.get("duration") or 0),
                int(record_id)))
    return {"ok": True, "turns": turns, "meta": meta2,
            "role_engine": role_engine, "fact_engine": fact_engine,
            "duration": tr.get("duration", 0)}


def delete_record(record_id: int, delete_files: bool = True) -> dict:
    r = db.query_one("SELECT * FROM intake_records WHERE id=?", (int(record_id),))
    if not r:
        return {"ok": False, "error": "记录不存在"}
    if delete_files:
        for key in ("audio_path", "docx_path", "transcript_path"):
            p = r[key]
            if p:
                try:
                    (VAULT_DIR / p).unlink(missing_ok=True)
                except Exception:
                    pass
    db.execute("DELETE FROM intake_records WHERE id=?", (int(record_id),))
    return {"ok": True}


def open_target(record_id: int, target: str) -> dict:
    r = db.query_one("SELECT * FROM intake_records WHERE id=?", (int(record_id),))
    if not r:
        return {"ok": False, "error": "记录不存在"}
    key = {"docx": "docx_path", "audio": "audio_path",
           "transcript": "transcript_path"}.get(target, "docx_path")
    rel = r[key]
    if not rel:
        return {"ok": False, "error": "文件尚未生成"}
    p = VAULT_DIR / rel
    if not p.exists():
        return {"ok": False, "error": "文件不存在，可能已被移动或删除"}
    return archive._open_with_system(p)

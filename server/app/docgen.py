# -*- coding: utf-8 -*-
"""空白法律文书生成：命名规则 = 两位序号 + 客户名称 + 案由 + 文件类型 + YYYYMMDD。
示例：00张先生劳动争议法律服务合同20260901.docx
"""
from __future__ import annotations
import hashlib
import re
from datetime import datetime
from pathlib import Path

from .config import DOCX_DIR
from . import db

INVALID = re.compile(r'[\\/:*?"<>|\s]+')


def safe_name(s: str) -> str:
    return INVALID.sub("", s or "")


def hash_file(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()


def next_seq(case_path: str) -> int:
    row = db.query_one("SELECT COALESCE(MAX(seq),-1) m FROM gen_docs WHERE case_path=?",
                       (case_path,))
    return int(row["m"]) + 1


def build_filename(seq: int, client: str, cause: str, doc_type: str,
                   when: datetime | None = None) -> str:
    when = when or datetime.now()
    return (f"{seq:02d}{safe_name(client)}{safe_name(cause)}"
            f"{safe_name(doc_type)}{when.strftime('%Y%m%d')}.docx")


def _set_cn_font(run, name="宋体", size=12):
    from docx.oxml.ns import qn
    run.font.name = name
    run.font.size = None
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def create_doc(case_path: str, client: str, cause: str, doc_type: str,
               basis: str = "", seq: int | None = None) -> dict:
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    if seq is None:
        seq = next_seq(case_path)
    filename = build_filename(seq, client, cause, doc_type)
    case_dir = DOCX_DIR / safe_name(f"{client}{cause}")
    case_dir.mkdir(parents=True, exist_ok=True)
    out = case_dir / filename

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "宋体"
    style.font.size = Pt(12)
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(doc_type)
    run.bold = True
    run.font.size = Pt(18)
    _set_cn_font(run, "宋体")
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run(f"委托人：{client}    案由：{cause}")
    sr.font.size = Pt(10.5)
    _set_cn_font(sr, "宋体")
    doc.add_paragraph("")
    for _ in range(18):                       # 预留空白正文行，即“空白文档”
        doc.add_paragraph("")
    foot = doc.add_paragraph()
    fr = foot.add_run(f"（本空白文书由本地工作台生成｜依据：{basis or '待补充'}）")
    fr.font.size = Pt(9)
    _set_cn_font(fr, "宋体")
    doc.save(out)

    rel_path = out.relative_to(DOCX_DIR.parent).as_posix()
    digest = hash_file(out)
    db.execute(
        "INSERT INTO gen_docs(case_path,seq,doc_type,filename,rel_path,created,system_hash,disk_hash,sync_state)"
        " VALUES(?,?,?,?,?,?,?,?,?)",
        (case_path, seq, doc_type, filename, rel_path,
         datetime.now().strftime("%Y-%m-%d %H:%M:%S"), digest, digest, "一致"))
    db.audit("生成文书", f"{filename} <- {case_path}")
    return {"filename": filename, "rel_path": rel_path, "seq": seq,
            "abs_path": str(out), "basis": basis}


def list_case_docs(case_path: str) -> list[dict]:
    return db.query("SELECT * FROM gen_docs WHERE case_path=? ORDER BY seq", (case_path,))


def generated_types(case_path: str) -> set[str]:
    return {r["doc_type"] for r in list_case_docs(case_path)}


def rescan_disk_state() -> list[dict]:
    """比对磁盘上的已生成文书与系统记录（需求12：本地文档被外部改动时提示）。"""
    changed = []
    from .config import VAULT_DIR
    for r in db.query("SELECT * FROM gen_docs"):
        p = VAULT_DIR / r["rel_path"]
        if not p.exists():
            if r["sync_state"] != "本地缺失":
                db.execute("UPDATE gen_docs SET sync_state='本地缺失' WHERE id=?", (r["id"],))
                changed.append({**r, "sync_state": "本地缺失"})
            continue
        disk = hash_file(p)
        if disk != r["system_hash"]:
            state = "外部已修改" if disk != r["disk_hash"] or r["sync_state"] == "外部已修改" else "外部已修改"
            if r["sync_state"] != "外部已修改":
                db.execute("UPDATE gen_docs SET disk_hash=?, sync_state='外部已修改' WHERE id=?",
                           (disk, r["id"]))
                changed.append({**r, "disk_hash": disk, "sync_state": "外部已修改"})
    return changed


def accept_sync(doc_id: int) -> dict:
    """用户确认后：以磁盘最新版本为准，同步系统记录。"""
    from .config import VAULT_DIR
    r = db.query_one("SELECT * FROM gen_docs WHERE id=?", (doc_id,))
    if not r:
        return {"ok": False, "error": "记录不存在"}
    p = VAULT_DIR / r["rel_path"]
    if not p.exists():
        return {"ok": False, "error": "本地文件不存在"}
    digest = hash_file(p)
    db.execute("UPDATE gen_docs SET system_hash=?, disk_hash=?, sync_state='一致' WHERE id=?",
               (digest, digest, doc_id))
    db.audit("同步文书", r["filename"])
    return {"ok": True}
